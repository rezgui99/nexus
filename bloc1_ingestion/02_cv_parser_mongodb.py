"""
NEXUS EMPLOI — Bloc 1 : Extraction de texte brut depuis un CV (PDF / DOCX)
===========================================================================
Auteur  : Aziz Rezgui (Projet de fin d'études MSc Data Management 2025-2026)
Objectif: Parser des CV dans différents formats et insérer le contenu brut
          dans MongoDB, prêts pour le traitement NER via LLM.

Stratégie multi-couches :
  1. unstructured  → meilleure extraction sémantique (éléments typés)
  2. PyPDF         → fallback léger pour les PDF simples
  3. python-docx   → fallback pour les fichiers Word .docx

Dépendances:
    pip install unstructured[pdf,docx] pypdf python-docx pymongo python-dotenv
    # Pour unstructured avec OCR (PDF scannés) :
    pip install "unstructured[pdf]" pdf2image pytesseract
"""

import os
import hashlib
import logging
from dataclasses import dataclass, asdict, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from pymongo import MongoClient, errors
from dotenv import load_dotenv

load_dotenv()

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s"
)
logger = logging.getLogger("nexus_cv_parser")

MONGO_URI  = os.getenv("MONGO_URI", "mongodb://localhost:27017")
DB_NAME    = "nexus_emploi"
COLLECTION = "raw_cv"

SUPPORTED_FORMATS = {".pdf", ".docx", ".doc", ".txt"}


# ---------------------------------------------------------------------------
# Schéma de données brutes — CV
# ---------------------------------------------------------------------------

@dataclass
class RawCV:
    """Représentation d'un CV brut avant traitement LLM."""
    # --- Identifiants ---
    fichier_nom:    str          # Nom du fichier original
    fichier_format: str          # pdf | docx | txt
    candidat_id:    str          # Hash du contenu (anonymisation par défaut)

    # --- Contenu brut ---
    texte_brut:     str          # Texte complet extrait
    elements:       list         # Éléments structurés (unstructured) ou [] si fallback

    # --- Métadonnées de collecte ---
    parsed_at:      str = field(default_factory=lambda: datetime.now(timezone.utc).isoformat())
    parser_utilise: str = "unknown"   # unstructured | pypdf | python-docx
    nb_pages:       Optional[int] = None
    nb_caracteres:  int = 0
    pipeline_version: str = "1.0.0"
    statut_traitement: str = "raw"    # raw → structured → matched
    content_hash:   str = ""

    def compute_hash(self) -> "RawCV":
        self.content_hash = hashlib.sha256(self.texte_brut.encode()).hexdigest()
        self.candidat_id  = self.content_hash[:16]   # Pseudo-anonymisation
        self.nb_caracteres = len(self.texte_brut)
        return self


# ---------------------------------------------------------------------------
# Parseurs
# ---------------------------------------------------------------------------

class UnstructuredParser:
    """
    Extraction via la librairie `unstructured`.
    Retourne à la fois le texte brut et une liste d'éléments typés
    (Title, NarrativeText, ListItem, Table…) — très utile pour le NER.
    """

    def parse(self, filepath: Path) -> tuple[str, list, str]:
        """
        Retourne (texte_brut, elements_structures, nom_parser).
        elements_structures : liste de dicts {type, text} utilisables par le LLM.
        """
        try:
            from unstructured.partition.auto import partition

            elements = partition(filename=str(filepath))
            texte_brut = "\n\n".join(el.text for el in elements if hasattr(el, "text") and el.text.strip())

            # Sérialisation légère des éléments pour MongoDB
            elements_json = [
                {
                    "type":     type(el).__name__,
                    "text":     el.text,
                    "metadata": el.metadata.to_dict() if hasattr(el.metadata, "to_dict") else {}
                }
                for el in elements if hasattr(el, "text")
            ]
            logger.info("✅ Unstructured — %d éléments extraits", len(elements))
            return texte_brut, elements_json, "unstructured"

        except ImportError:
            logger.warning("⚠️  unstructured non installé — fallback vers PyPDF/python-docx")
            return "", [], "fallback"
        except Exception as e:
            logger.warning("⚠️  Erreur unstructured (%s) — fallback activé", e)
            return "", [], "fallback"


class PyPDFParser:
    """Extraction légère pour PDF. Fallback si unstructured échoue."""

    def parse(self, filepath: Path) -> tuple[str, int]:
        """Retourne (texte_brut, nb_pages)."""
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(filepath))
            pages_text = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    pages_text.append(f"--- Page {i+1} ---\n{text.strip()}")

            texte = "\n\n".join(pages_text)
            logger.info("✅ PyPDF — %d pages, %d caractères", len(reader.pages), len(texte))
            return texte, len(reader.pages)

        except ImportError:
            logger.error("❌ pypdf non installé : pip install pypdf")
            return "", 0
        except Exception as e:
            logger.error("❌ Erreur PyPDF : %s", e)
            return "", 0


class DocxParser:
    """Extraction pour fichiers Word .docx."""

    def parse(self, filepath: Path) -> tuple[str, list]:
        """Retourne (texte_brut, elements_structures)."""
        try:
            from docx import Document

            doc = Document(str(filepath))
            elements = []
            paragraphes = []

            for para in doc.paragraphs:
                if para.text.strip():
                    # Conserver le style (Heading 1, Heading 2, Normal…)
                    elements.append({
                        "type": para.style.name,
                        "text": para.text.strip(),
                        "metadata": {}
                    })
                    paragraphes.append(para.text.strip())

            # Extraction des tableaux (compétences, expériences en tableau)
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        elements.append({"type": "Table", "text": row_text, "metadata": {}})
                        paragraphes.append(row_text)

            texte = "\n\n".join(paragraphes)
            logger.info("✅ python-docx — %d éléments extraits", len(elements))
            return texte, elements

        except ImportError:
            logger.error("❌ python-docx non installé : pip install python-docx")
            return "", []
        except Exception as e:
            logger.error("❌ Erreur python-docx : %s", e)
            return "", []


# ---------------------------------------------------------------------------
# Orchestrateur principal
# ---------------------------------------------------------------------------

class CVParser:
    """
    Orchestre les différents parseurs avec une stratégie de fallback.
    Priorité : unstructured → PyPDF/DocxParser → erreur.
    """

    def __init__(self):
        self.unstructured = UnstructuredParser()
        self.pypdf        = PyPDFParser()
        self.docx         = DocxParser()

    def parse(self, filepath: str | Path) -> Optional[RawCV]:
        """
        Parse un CV et retourne un RawCV prêt pour MongoDB.
        Gère les fallbacks automatiquement selon le format.
        """
        filepath = Path(filepath)

        if not filepath.exists():
            logger.error("❌ Fichier introuvable : %s", filepath)
            return None

        suffix = filepath.suffix.lower()
        if suffix not in SUPPORTED_FORMATS:
            logger.error("❌ Format non supporté : %s", suffix)
            return None

        logger.info("📄 Parsing du CV : %s", filepath.name)

        texte, elements, parser_name, nb_pages = "", [], "unknown", None

        # ---- Tentative 1 : unstructured (tous formats) ----
        texte, elements, parser_name = self.unstructured.parse(filepath)

        # ---- Tentative 2 : fallback spécifique par format ----
        if not texte.strip():
            if suffix == ".pdf":
                texte, nb_pages = self.pypdf.parse(filepath)
                parser_name = "pypdf"
                elements = []
            elif suffix in (".docx", ".doc"):
                texte, elements = self.docx.parse(filepath)
                parser_name = "python-docx"
            elif suffix == ".txt":
                texte = filepath.read_text(encoding="utf-8", errors="replace")
                parser_name = "txt-native"

        if not texte.strip():
            logger.error("❌ Impossible d'extraire le texte de %s", filepath.name)
            return None

        cv = RawCV(
            fichier_nom    = filepath.name,
            fichier_format = suffix.lstrip("."),
            candidat_id    = "",                  # Sera calculé par compute_hash()
            texte_brut     = texte,
            elements       = elements,
            parser_utilise = parser_name,
            nb_pages       = nb_pages,
        ).compute_hash()

        logger.info(
            "✅ CV parsé — %s | Parser: %s | %d caractères",
            cv.fichier_nom, cv.parser_utilise, cv.nb_caracteres
        )
        return cv


# ---------------------------------------------------------------------------
# Stockage MongoDB
# ---------------------------------------------------------------------------

class CVMongoInserter:
    """Gère l'insertion des CV bruts dans MongoDB."""

    def __init__(self, uri: str, db: str, collection: str):
        self.client = MongoClient(uri, serverSelectionTimeoutMS=5000)
        self.col = self.client[db][collection]
        self.col.create_index("content_hash", unique=True, sparse=True)
        self.col.create_index("statut_traitement")
        logger.info("✅ MongoDB prêt — %s.%s", db, collection)

    def insert(self, cv: RawCV) -> Optional[str]:
        doc = asdict(cv)
        result = self.col.update_one(
            {"content_hash": cv.content_hash},
            {"$setOnInsert": doc},
            upsert=True
        )
        if result.upserted_id:
            logger.info("✅ CV inséré — id: %s | %s", result.upserted_id, cv.fichier_nom)
            return str(result.upserted_id)
        logger.debug("⏩ CV déjà présent — hash: %s", cv.content_hash[:12])
        return None

    def close(self):
        self.client.close()


# ---------------------------------------------------------------------------
# Pipeline
# ---------------------------------------------------------------------------

def ingest_cv(cv_path: str) -> Optional[dict]:
    """
    Fonction de haut niveau : parse un CV et l'insère dans MongoDB.
    Retourne le document inséré ou None.
    """
    parser  = CVParser()
    cv      = parser.parse(cv_path)

    if not cv:
        return None

    inserter = CVMongoInserter(MONGO_URI, DB_NAME, COLLECTION)
    inserted_id = inserter.insert(cv)
    inserter.close()

    if inserted_id:
        return {"_id": inserted_id, **asdict(cv)}
    return None


def ingest_dossier(dossier: str) -> list[dict]:
    """
    Ingère tous les CV d'un répertoire.
    Utile pour traiter un corpus de candidats en batch.
    """
    parser   = CVParser()
    inserter = CVMongoInserter(MONGO_URI, DB_NAME, COLLECTION)
    resultats = []

    for fichier in Path(dossier).iterdir():
        if fichier.suffix.lower() in SUPPORTED_FORMATS:
            cv = parser.parse(fichier)
            if cv:
                inserted_id = inserter.insert(cv)
                if inserted_id:
                    resultats.append({"fichier": fichier.name, "_id": inserted_id})

    inserter.close()
    logger.info("🏁 %d CV ingérés depuis %s", len(resultats), dossier)
    return resultats


# ---------------------------------------------------------------------------
# Point d'entrée
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import sys

    if len(sys.argv) > 1:
        path = sys.argv[1]
        if Path(path).is_dir():
            ingest_dossier(path)
        else:
            ingest_cv(path)
    else:
        print("Usage :")
        print("  python 02_cv_parser_mongodb.py mon_cv.pdf")
        print("  python 02_cv_parser_mongodb.py /dossier/cv/")